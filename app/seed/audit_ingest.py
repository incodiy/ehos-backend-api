"""Audit Excel ingestion: audit_sessions + audit_item_scores.

Reads the 2026 auditor checklist workbooks (GM / Housekeeping / Kitchen & FB /
Security Risk Management) and maps each file into a ChecklistTemplate snapshot
+ one AuditSession per (hotel, department, period, audit_type) with item-level
scores.

Formats supported (mapper per format):
- TRAFFIC_LIGHT (GM, HK, SRM): YES(90) / NO(0) / NEED REVIEW(45) columns.
- BINARY_COUNT  (Kitchen & FB): per-item YES/NO → 1/0 point.

Deterministic & idempotent (Constraint H1-H4): re-running replaces the existing
session + scores for the same unique key. Follow-up workbooks (findings-style
sheets) are intentionally deferred to the CAPA/findings phase.
"""

import os
import re
import uuid
from collections import Counter
from dataclasses import dataclass, field
from datetime import UTC, date, datetime

import openpyxl
from sqlalchemy import delete, select
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import (
    AuditItemScore,
    AuditSession,
    ChecklistItem,
    ChecklistSection,
    ChecklistTemplate,
    Hotel,
)

MONTH_INDEX = {
    name.lower(): idx
    for idx, name in enumerate(
        [
            "January",
            "February",
            "March",
            "April",
            "May",
            "June",
            "July",
            "August",
            "September",
            "October",
            "November",
            "December",
        ],
        start=1,
    )
}

DATE_PATTERN = re.compile(r"(\d{1,2})\s*(?:[-–—]\s*(\d{1,2}))?\s+([A-Z][a-z]+)\s+(\d{4})")

FAMILY_GM = "GM"
FAMILY_HK = "HOUSEKEEPING"
FAMILY_KFB = "KITCHEN_FB"
FAMILY_SRM = "SECURITY_RISK"

TEMPLATE_NAMES = {
    FAMILY_GM: "General Manager Audit Checklist",
    FAMILY_HK: "Housekeeping Audit Checklist",
    FAMILY_KFB: "Kitchen & F&B Audit Checklist",
    FAMILY_SRM: "Security Risk Management Audit Checklist",
}

TEMPLATE_VERSION = "2026"

TOTAL_LABELS = {"POINTS", "SCORE IN %", "TOTAL", "TOTAL POSSIBLE", "POINTS ACHIEVED", "TOTAL POSSIBLE SCORE"}


@dataclass
class ParsedItem:
    section_code: str
    section_name: str
    code: str
    question: str
    value: str | None
    score: float
    note: str | None


@dataclass
class ParsedFile:
    hotel_code: str
    department: str
    audit_type: str
    date_start: date
    date_end: date
    sections: list[tuple[str, str]] = field(default_factory=list)
    items: list[ParsedItem] = field(default_factory=list)


def _numcode(value: object) -> bool:
    if value is None:
        return False
    if isinstance(value, (int, float)):
        return True
    return bool(re.fullmatch(r"\d+(\.\d+)?", str(value).strip()))


def _is_section_letter(value: object) -> bool:
    if not isinstance(value, str):
        return False
    return len(value.strip()) == 1 and value.strip().isalpha()


def _clean(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def _audit_date(wb: openpyxl.Workbook) -> tuple[date, date]:
    """Locate the audit date range, preferring Summary/Follow-up sheets."""
    best: tuple[int, date, date] | None = None
    for sheet_name in wb.sheetnames:
        upper = sheet_name.upper()
        bonus = 8 if ("SUMMARY" in upper or "FOLLOW UP" in upper) else 0
        ws = wb[sheet_name]
        for raw in ws.iter_rows(values_only=True):
            for cell in raw:
                if not isinstance(cell, str):
                    continue
                for match in DATE_PATTERN.finditer(cell):
                    month = MONTH_INDEX.get(match.group(3).lower())
                    if month is None:
                        continue
                    day1 = int(match.group(1))
                    day2 = int(match.group(2)) if match.group(2) else day1
                    year = int(match.group(4))
                    if not (1 <= day1 <= 31 and 1 <= day2 <= 31 and 2020 <= year <= 2027):
                        continue
                    score = bonus
                    score += 5 if match.group(2) else 0
                    score += 3 if year == 2026 else 0
                    if "by " in cell[: match.start()].lower()[-30:]:
                        score -= 4
                    candidate = (score, date(year, month, day1), date(year, month, day2))
                    if best is None or candidate[0] > best[0]:
                        best = candidate
    if best is None:
        raise ValueError(f"No audit date found in workbook ({wb.sheetnames})")
    return best[1], best[2]


def _hotel_code(path: str, known_codes: set[str]) -> str:
    upper = os.path.basename(path).upper()
    for code in sorted(known_codes, key=len, reverse=True):
        if code in upper:
            return code
    raise ValueError(f"Hotel code not found in filename: {os.path.basename(path)}")


def _file_kind(path: str) -> tuple[str, str]:
    upper = os.path.basename(path).upper()
    is_follow_up = "FOLLOW UP" in upper
    if "GENERAL MANAGER" in upper:
        return FAMILY_GM, ""
    if "HOUSEKEEPING" in upper:
        return FAMILY_HK, "FOLLOW_UP" if is_follow_up else "FULL"
    if "KITCHEN" in upper or "F&B" in upper:
        return FAMILY_KFB, "FOLLOW_UP" if is_follow_up else "FULL"
    if "SECURITY" in upper:
        return FAMILY_SRM, "FOLLOW_UP" if is_follow_up else "FULL"
    raise ValueError(f"Unrecognised audit file: {os.path.basename(path)}")


def parse_workbook(path: str, known_codes: set[str]) -> ParsedFile:
    department, kind = _file_kind(path)
    hotel_code = _hotel_code(path, known_codes)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        date_start, date_end = _audit_date(wb)

        if department == FAMILY_GM:
            parsed = _parse_blocks(wb, "Sheet1", traffic=True)
        elif department in (FAMILY_HK, FAMILY_SRM):
            sheet = _key_sheet(wb, ("HK Audit", "Security Risk Management Audit"))
            parsed = _parse_blocks(wb, sheet, traffic=True)
        else:
            sheet = _key_sheet(wb, ("F&B AUDIT CHECKLIST", "F&B AUDIT FOLLOW UP"))
            parsed = _parse_kfb(wb, sheet)
        sections, items = parsed
    finally:
        wb.close()

    return ParsedFile(
        hotel_code=hotel_code,
        department=department,
        audit_type=kind or "FULL",
        date_start=date_start,
        date_end=date_end,
        sections=sections,
        items=items,
    )


def _key_sheet(wb: openpyxl.Workbook, candidates: tuple[str, ...]) -> str:
    for candidate in candidates:
        if candidate in wb.sheetnames:
            return candidate
    raise ValueError(f"Checklist sheet not found (have {wb.sheetnames})")


def _is_header(raw: tuple) -> bool:
    resp = raw[2:5]
    if len(resp) < 3 or not isinstance(resp[0], str):
        return False
    return "YES" in resp[0].upper() and "NEED" in str(resp[2]).upper()


def _parse_blocks(wb: openpyxl.Workbook, sheet: str, traffic: bool) -> tuple[list[tuple[str, str]], list[ParsedItem]]:
    ws = wb[sheet]
    sections: list[tuple[str, str]] = []
    items: list[ParsedItem] = []
    section_code: str | None = None
    section_name: str | None = None
    seq = 0
    started = False

    for raw in ws.iter_rows(values_only=True):
        col_a, col_b = raw[0], raw[1]
        resp = raw[2:5]
        comment = raw[5] if len(raw) > 5 else None

        if not started:
            if _is_header(raw):
                started = True
            else:
                continue

        title = _clean(col_b)
        if title and title.upper() in TOTAL_LABELS:
            continue
        if col_a is None and title:
            seq += 1
            section_code, section_name = f"S{seq:02d}", title
            sections.append((section_code, section_name))
            continue
        if _is_section_letter(col_a) and title and not _is_numeric(resp[0]):
            section_code, section_name = _clean(col_a) or f"S{seq + 1:02d}", title
            if not any(code == section_code for code, _ in sections):
                sections.append((section_code, section_name))
            continue
        if section_code is None:
            continue

        if _numcode(col_a) and title:
            value, score, note = _traffic_response(resp, comment)
            items.append(
                ParsedItem(
                    section_code=section_code,
                    section_name=section_name,
                    code=str(col_a),
                    question=title,
                    value=value,
                    score=score,
                    note=note,
                )
            )
    return sections, items


def _is_numeric(value: object) -> bool:
    if isinstance(value, (int, float)):
        return True
    return isinstance(value, str) and bool(re.fullmatch(r"\d+(\.\d+)?", value.strip()))


def _traffic_response(resp: tuple, comment) -> tuple[str | None, float, str | None]:
    buckets = ((resp[0], "YES", 90.0), (resp[1], "NO", 0.0), (resp[2], "NEED REVIEW", 45.0))
    noted = [(label, pts) for cell, label, pts in buckets if _is_numeric(cell)]
    if not noted:
        return None, 0.0, _clean(comment)
    label, pts = noted[0]
    return label, pts, _clean(comment)


def _parse_kfb(wb: openpyxl.Workbook, sheet: str) -> tuple[list[tuple[str, str]], list[ParsedItem]]:
    ws = wb[sheet]
    sections: list[tuple[str, str]] = []
    items: list[ParsedItem] = []
    section_code: str | None = None
    section_name: str | None = None

    for raw in ws.iter_rows(values_only=True):
        col_a, col_b = raw[0], raw[1]
        resp = raw[4] if len(raw) > 4 else None
        comment = raw[5] if len(raw) > 5 else None

        title = _clean(col_b)
        if title and title.upper() in TOTAL_LABELS:
            continue
        if _is_section_letter(col_a) and title and (resp is None or "YES" in str(resp).upper()):
            section_code, section_name = _clean(col_a), title
            sections.append((section_code, section_name))
            continue
        if section_code is None or not _numcode(col_a) or not title:
            continue

        value = resp.upper() if resp is not None else None
        if value not in ("YES", "NO"):
            value = None
        score = 1.0 if value == "YES" else 0.0
        items.append(
            ParsedItem(
                section_code=section_code,
                section_name=section_name,
                code=str(col_a),
                question=title,
                value=value,
                score=score,
                note=_clean(comment),
            )
        )
    return sections, items


def resolve_files(audit_dir: str, known_codes: set[str]) -> list[str]:
    """Discover workbook paths, dropping (1) duplicate copies when a base exists."""
    paths = sorted(f for f in os.listdir(audit_dir) if f.lower().endswith(".xlsx"))
    by_key: dict[tuple[str, str], str] = {}
    for filename in paths:
        try:
            department, _ = _file_kind(filename)
        except ValueError:
            continue
        hotel = _hotel_code(filename, known_codes)
        if "FOLLOW UP" in filename.upper():
            continue
        key = (department, hotel)
        if key not in by_key or "(1)" not in filename:
            by_key[key] = filename
    return [os.path.join(audit_dir, by_key[key]) for key in sorted(by_key)]


async def seed_audit_ingest(session: AsyncSession, audit_dir: str, system_user_id: uuid.UUID) -> dict:
    hotel_codes = set((await session.execute(select(Hotel.code))).scalars().all())
    paths = resolve_files(audit_dir, hotel_codes)

    template_cache: dict[str, dict] = {}
    result: dict[str, int] = Counter()

    for path in paths:
        parsed = parse_workbook(path, hotel_codes)
        template_id = await _ensure_template(session, parsed.department, system_user_id, template_cache)
        item_ids = await _ensure_sections(session, template_id, parsed.department, parsed.sections, parsed.items)

        session_key = parsed.hotel_code
        await _replace_session(
            session,
            parsed,
            template_id,
            item_ids,
            system_user_id,
            hotel_codes,
        )
        result[session_key] += 1

    return dict(result)


async def _ensure_template(
    session: AsyncSession, department: str, published_by: uuid.UUID, cache: dict[str, dict]
) -> uuid.UUID:
    if department in cache:
        return cache[department]["id"]
    name = TEMPLATE_NAMES[department]
    template = await session.scalar(
        select(ChecklistTemplate).where(
            ChecklistTemplate.department == department,
            ChecklistTemplate.name == name,
            ChecklistTemplate.version == TEMPLATE_VERSION,
        )
    )
    if template is None:
        template = ChecklistTemplate(
            department=department,
            name=name,
            version=TEMPLATE_VERSION,
            status="LOCKED",
            published_by=published_by,
        )
        session.add(template)
        await session.flush()
    cache[department] = {"id": template.id, "sections": {}}
    return template.id


async def _ensure_sections(
    session: AsyncSession,
    template_id: uuid.UUID,
    department: str,
    sections: list[tuple[str, str]],
    items: list[ParsedItem],
) -> dict[str, uuid.UUID]:
    existing = dict(
        (
            await session.execute(
                select(ChecklistSection.code, ChecklistSection.id).where(ChecklistSection.template_id == template_id)
            )
        ).all()
    )
    known_items = {
        (section_id, code): item_id
        for section_id, code, item_id in (
            await session.execute(
                select(ChecklistItem.section_id, ChecklistItem.code, ChecklistItem.id).where(
                    ChecklistItem.section_id.in_(list(existing.values()))
                )
            )
        ).all()
    }
    is_traffic = department != FAMILY_KFB

    for code, name in sections:
        if code not in existing:
            row = ChecklistSection(template_id=template_id, code=code, name=name, sort_order=len(existing))
            session.add(row)
            await session.flush()
            existing[code] = row.id

    item_ids: dict[str, uuid.UUID] = {}
    for index, item in enumerate(items):
        section_id = existing[item.section_code]
        item_id = known_items.get((section_id, item.code))
        if item_id is None:
            item_id = await session.scalar(
                select(ChecklistItem.id).where(
                    ChecklistItem.section_id == section_id,
                    ChecklistItem.code == item.code,
                )
            )
        if item_id is None:
            rubric = "TRAFFIC_LIGHT" if is_traffic else "BINARY_COUNT"
            max_score = 90.0 if is_traffic else 1.0
            row = ChecklistItem(
                section_id=section_id,
                code=item.code,
                question_text=item.question,
                rubric_type=rubric,
                max_score=max_score,
                weight=1.0,
                na_allowed=False,
                is_life_safety=False,
                sort_order=index,
            )
            session.add(row)
            await session.flush()
            item_id = row.id
        item_ids[f"{item.section_code}:{item.code}"] = item_id

    return item_ids


async def _replace_session(
    session: AsyncSession,
    parsed: ParsedFile,
    template_id: uuid.UUID,
    item_ids: dict[str, uuid.UUID],
    scored_by: uuid.UUID,
    hotel_codes: set[str],
) -> None:
    hotel_id = await session.scalar(select(Hotel.id).where(Hotel.code == parsed.hotel_code))

    existing = await session.scalar(
        select(AuditSession.id).where(
            AuditSession.hotel_id == hotel_id,
            AuditSession.department == parsed.department,
            AuditSession.date_start == parsed.date_start,
            AuditSession.audit_type == parsed.audit_type,
        )
    )
    if existing is not None:
        await session.execute(delete(AuditItemScore).where(AuditItemScore.session_id == existing))
        await session.execute(delete(AuditSession).where(AuditSession.id == existing))

    total = sum(item.score for item in parsed.items)
    count = len(parsed.items)
    if parsed.department == FAMILY_KFB:
        max_score = float(count)
    else:
        max_score = 90.0 * count
    pct = (total / max_score) if max_score else 0.0

    per_section: dict[str, dict] = {}
    for item in parsed.items:
        bucket = per_section.setdefault(item.section_code, {"score": 0.0, "count": 0})
        bucket["score"] += item.score
        bucket["count"] += 1
    subcategory_scores = {
        code: {
            "score": bucket["score"],
            "max": ((bucket["count"] * 90.0) if parsed.department != FAMILY_KFB else float(bucket["count"])),
            "count": bucket["count"],
            "pct": round(
                (bucket["score"] / (bucket["count"] * 90.0 if parsed.department != FAMILY_KFB else bucket["count"]))
                * 100,
                2,
            ),
        }
        for code, bucket in per_section.items()
    }

    pct_score = round(pct * 100, 2)
    row = AuditSession(
        hotel_id=hotel_id,
        template_id=template_id,
        department=parsed.department,
        audit_type=parsed.audit_type,
        status="PUBLISHED",
        auditor_id=scored_by,
        date_start=parsed.date_start,
        date_end=parsed.date_end,
        published_at=datetime.now(UTC),
        pass_fail="PASS" if pct >= 0.8 else "FAIL",
        total_score=pct_score,
        department_breakdown={
            parsed.department: {"score": pct_score, "max": 100.0, "pct": round(pct, 4), "items": count}
        },
        subcategory_scores=subcategory_scores,
        origin="SYSTEM",
        created_by=scored_by,
        updated_by=scored_by,
    )
    session.add(row)
    await session.flush()

    score_rows = []
    for item in parsed.items:
        score_rows.append(
            {
                "session_id": row.id,
                "item_id": item_ids[f"{item.section_code}:{item.code}"],
                "room_ref": None,
                "value": item.value,
                "score": item.score,
                "is_na": False,
                "note": item.note,
                "scored_by": scored_by,
                "scored_at": datetime.now(UTC),
                "updated_at": datetime.now(UTC),
            }
        )
    if score_rows:
        await session.execute(pg_insert(AuditItemScore).values(score_rows))
