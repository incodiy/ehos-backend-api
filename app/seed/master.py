"""Master-data seeder: brands, provinces, regions, hotels, hotel_departments.

Deterministic & idempotent (Constraint H1-H4). Source of truth: `crm/Master Data Hotel.xlsx`.
"""

from collections import Counter
from datetime import date, datetime, timezone

import openpyxl
from geoalchemy2.elements import WKTElement
from sqlalchemy import select, text
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Brand, Hotel, HotelDepartment, Province, Region
from app.seed.data import (
    BRAND_BY_SOURCE_LABEL,
    BRANDS,
    CITY_PROVINCE,
    COORDINATE_FALLBACK,
    HOTEL_DEPARTMENTS,
    HOTEL_STATUS_MAP,
    PROVINCES,
    REGION_CODE,
)

SOURCE_FILE = "Master Data Hotel.xlsx"


def _norm_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value or "").strip()
    if not raw or raw == "1000-01-01":
        return None
    try:
        return datetime.fromisoformat(raw).date()
    except ValueError:
        return None


async def seed_brands(session: AsyncSession) -> None:
    stmt = pg_insert(Brand).values(BRANDS)
    stmt = stmt.on_conflict_do_update(
        index_elements=[Brand.code],
        set_={"name": stmt.excluded.name, "tier": stmt.excluded.tier, "status": stmt.excluded.status},
    )
    await session.execute(stmt)

    # Constraint H3: Skenario soft-delete untuk pengujian isolasi query
    soft_deleted_code = "TEST_LEGACY_BRAND"
    existing_soft = await session.scalar(select(Brand).where(Brand.code == soft_deleted_code))
    if not existing_soft:
        brand_sd = Brand(
            code=soft_deleted_code,
            name="Brand Prototype Historical Archive",
            tier="Midscale",
            status="RETIRED",
        )
        brand_sd.deleted_at = datetime.now(timezone.utc)
        session.add(brand_sd)
        await session.flush()


async def seed_provinces(session: AsyncSession) -> None:
    stmt = pg_insert(Province).values(PROVINCES)
    stmt = stmt.on_conflict_do_update(
        index_elements=[Province.code],
        set_={"name": stmt.excluded.name},
    )
    await session.execute(stmt)


async def _load_hotel_rows(data_dir: str) -> tuple[list[dict], list[dict]]:
    path = f"{data_dir}/{SOURCE_FILE}"
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["md_hotel"]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    col = {name: i for i, name in enumerate(header)}

    region_sales: dict[str, Counter] = {}
    rows: list[dict] = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        code = raw[col["Hotel Code"]]
        if not code:
            continue
        brand_label = raw[col["Brand"]]
        brand_code = BRAND_BY_SOURCE_LABEL.get(brand_label)
        if brand_code is None:
            raise ValueError(f"Unknown brand label `{brand_label}` for hotel {code}")
        region_name = raw[col["Region"]]
        city = raw[col["City"]]
        if region_name not in REGION_CODE:
            raise ValueError(f"Unknown region `{region_name}` for hotel {code}")
        if city not in CITY_PROVINCE:
            raise ValueError(f"Unknown city `{city}` for hotel {code}")
        region_sales.setdefault(region_name, Counter())[raw[col["Sales Region"]]] += 1
        rows.append(
            {
                "code": code,
                "name": (raw[col["Hotel Name"]] or "").strip().rstrip(",").strip(),
                "brand_code": brand_code,
                "region_code": REGION_CODE[region_name],
                "province_code": CITY_PROVINCE[city],
                "city": city,
                "opening_date": _norm_date(raw[col["Opening Date"]]),
                "terminate_date": _norm_date(raw[col["Terminate Date"]]),
                "status": HOTEL_STATUS_MAP.get(raw[col["Status"]], "ACTIVE"),
                "lat": raw[col["Latitude"]],
                "lon": raw[col["Longitude"]],
            }
        )
    wb.close()
    rows.sort(key=lambda x: x["code"])

    regions = [
        {
            "code": REGION_CODE[name],
            "name": name,
            "country": "Indonesia",
            "sales_region": region_sales[name].most_common(1)[0][0],
            "status": "ACTIVE",
        }
        for name in REGION_CODE
    ]

    # Constraint H1-H4 Multi-scenario simulations:
    # 1. Happy path (14 active regions linked to 106 hotels & ROMs)
    # 2. Edge case 1: Region without hotels (zero-hotel empty state test)
    # 3. Edge case 2: Inactive region (unit reorganization/transition)
    # 4. Edge case 3: Retired legacy region (historical 2024 boundary)
    # 5. Edge case 4: Soft-deleted test region (verify deleted_at isolation)
    # 6. Edge case 5: International multi-country regional scope (Vietnam/APAC)
    extra_scenarios = [
        {
            "code": "NUSANTARA",
            "name": "Nusantara Capital City (IKN)",
            "country": "Indonesia",
            "sales_region": "East Kalimantan & Nusantara Sales",
            "status": "ACTIVE",
        },
        {
            "code": "TIMOR_BARAT",
            "name": "West Timor Special Region",
            "country": "Indonesia",
            "sales_region": "Sunda Kecil Sales",
            "status": "INACTIVE",
        },
        {
            "code": "MALUKU_PAPUA",
            "name": "Maluku & Papua Combined (Legacy 2024)",
            "country": "Indonesia",
            "sales_region": "East Indonesia Combined",
            "status": "RETIRED",
        },
        {
            "code": "TEST_EXPANSION",
            "name": "Test Regional Expansion Unit",
            "country": "Indonesia",
            "sales_region": "R&D Expansion",
            "status": "ACTIVE",
        },
        {
            "code": "APAC_OVERSEAS",
            "name": "Indochina & Philippines",
            "country": "Vietnam",
            "sales_region": "APAC International Division",
            "status": "ACTIVE",
        },
    ]
    regions.extend(extra_scenarios)
    regions.sort(key=lambda x: x["code"])
    return rows, regions


async def seed_regions_and_hotels(session: AsyncSession, data_dir: str) -> None:
    rows, regions = await _load_hotel_rows(data_dir)

    if regions:
        stmt = pg_insert(Region).values(regions)
        stmt = stmt.on_conflict_do_update(
            index_elements=[Region.code],
            set_={
                "name": stmt.excluded.name,
                "country": stmt.excluded.country,
                "sales_region": stmt.excluded.sales_region,
                "status": stmt.excluded.status,
            },
        )
        await session.execute(stmt)

        # Soft-delete test region for query verification
        await session.execute(
            text("UPDATE regions SET deleted_at = NOW() WHERE code = 'TEST_EXPANSION' AND deleted_at IS NULL")
        )

    brand_ids = dict((await session.execute(select(Brand.code, Brand.id))).all())
    region_ids = dict((await session.execute(select(Region.code, Region.id))).all())
    province_ids = dict((await session.execute(select(Province.code, Province.id))).all())

    hotel_values = []
    for row in rows:
        lat, lon = row["lat"], row["lon"]
        if lat is None or lon is None:
            lat, lon = COORDINATE_FALLBACK[row["code"]]

        # Skenario multi-status FSM: TEMPORARILY_CLOSED untuk hotel dalam masa renovasi besar
        status = row["status"]
        if row["code"] in ("ATTB", "LHSB", "HCC"):
            status = "TEMPORARILY_CLOSED"

        # Simulasi MICE facilities komprehensif (PRD-F-01 & Constraint H1-H4)
        mice_facilities = None
        if row["brand_code"] in ("SBH", "SBR", "GSB", "SBEC", "MDK") or row["code"] in ("CWS", "SBAI", "SQYO"):
            base_pax = 500 + (sum(ord(c) for c in row["code"]) % 8) * 100
            rooms = 4 + (sum(ord(c) for c in row["code"]) % 6)
            mice_facilities = {
                "ballroom_capacity": base_pax,
                "meeting_rooms": rooms,
                "has_videotron": (base_pax >= 800),
            }

        # Variasi geofence radius (resort luas 450m vs city hotel 150-300m)
        radius = 200
        if "resort" in row["name"].lower() or row["brand_code"] == "SBR":
            radius = 450
        elif row["brand_code"] in ("SBH", "GSB"):
            radius = 300
        elif row["brand_code"] in ("ZES", "SBEX"):
            radius = 150

        hotel_values.append(
            {
                "code": row["code"],
                "name": row["name"],
                "brand_id": brand_ids[row["brand_code"]],
                "region_id": region_ids[row["region_code"]],
                "province_id": province_ids[row["province_code"]],
                "city": row["city"],
                "geo": WKTElement(f"SRID=4326;POINT({lon} {lat})"),
                "geofence_radius_meters": radius,
                "mice_facilities": mice_facilities,
                "opening_date": row["opening_date"],
                "terminate_date": row["terminate_date"],
                "status": status,
            }
        )

    if hotel_values:
        stmt = pg_insert(Hotel).values(hotel_values)
        stmt = stmt.on_conflict_do_update(
            index_elements=[Hotel.code],
            set_={
                "name": stmt.excluded.name,
                "brand_id": stmt.excluded.brand_id,
                "region_id": stmt.excluded.region_id,
                "province_id": stmt.excluded.province_id,
                "city": stmt.excluded.city,
                "geo": stmt.excluded.geo,
                "geofence_radius_meters": stmt.excluded.geofence_radius_meters,
                "mice_facilities": stmt.excluded.mice_facilities,
                "opening_date": stmt.excluded.opening_date,
                "terminate_date": stmt.excluded.terminate_date,
                "status": stmt.excluded.status,
            },
        )
        await session.execute(stmt)

    hotel_ids = dict((await session.execute(select(Hotel.code, Hotel.id))).all())
    dept_values = [
        {
            "hotel_id": hotel_ids[row["code"]],
            "code": dept["code"],
            "name": dept["name"],
        }
        for row in rows
        for dept in HOTEL_DEPARTMENTS
    ]
    if dept_values:
        stmt = pg_insert(HotelDepartment).values(dept_values)
        stmt = stmt.on_conflict_do_update(
            index_elements=[HotelDepartment.hotel_id, HotelDepartment.code],
            set_={"name": stmt.excluded.name},
        )
        await session.execute(stmt)


async def seed_master(session: AsyncSession, data_dir: str) -> None:
    await seed_brands(session)
    await seed_provinces(session)
    await seed_regions_and_hotels(session, data_dir)
